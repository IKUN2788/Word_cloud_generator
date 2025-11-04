import tkinter as tk
from tkinter import ttk, filedialog, messagebox, colorchooser
import jieba
from wordcloud import WordCloud
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from PIL import Image, ImageTk
import numpy as np
import io
from collections import Counter
import re
import os

class WordCloudGenerator:
    def __init__(self, root):
        self.root = root
        self.root.title("词云图生成器")
        self.root.geometry("1200x860")
        self.root.configure(bg='#f0f0f0')
        
        # 初始化变量
        self.text_content = ""
        self.current_wordcloud = None
        self.current_figure = None
        
        # 创建界面
        self.create_widgets()
        
    def create_widgets(self):
        # 主标题
        title_frame = tk.Frame(self.root, bg='#f0f0f0')
        title_frame.pack(fill='x', padx=10, pady=5)
        
        title_label = tk.Label(title_frame, text="☁️ 词云图生成器", 
                              font=('Arial', 20, 'bold'), bg='#f0f0f0', fg='#333')
        title_label.pack()
        
        # 创建主框架
        main_frame = tk.Frame(self.root, bg='#f0f0f0')
        main_frame.pack(fill='both', expand=True, padx=10, pady=5)
        
        # 左侧面板 - 输入和设置
        left_frame = tk.Frame(main_frame, bg='white', relief='raised', bd=1)
        left_frame.pack(side='left', fill='both', expand=False, padx=(0, 5))
        left_frame.configure(width=400)
        
        # 右侧面板 - 显示区域
        right_frame = tk.Frame(main_frame, bg='white', relief='raised', bd=1)
        right_frame.pack(side='right', fill='both', expand=True, padx=(5, 0))
        
        self.create_left_panel(left_frame)
        self.create_right_panel(right_frame)
        
    def create_left_panel(self, parent):
        # 文本输入区域
        input_frame = tk.LabelFrame(parent, text="📝 文本输入", font=('Arial', 12, 'bold'), 
                                   bg='white', fg='#333', padx=10, pady=10)
        input_frame.pack(fill='x', padx=10, pady=5)
        
        # 输入方式选择
        self.input_method = tk.StringVar(value="直接输入")
        method_frame = tk.Frame(input_frame, bg='white')
        method_frame.pack(fill='x', pady=5)
        
        tk.Radiobutton(method_frame, text="直接输入/粘贴文本", variable=self.input_method, 
                      value="直接输入", bg='white', command=self.toggle_input_method).pack(anchor='w')
        tk.Radiobutton(method_frame, text="上传文件", variable=self.input_method, 
                      value="上传文件", bg='white', command=self.toggle_input_method).pack(anchor='w')
        
        # 文本输入框
        self.text_frame = tk.Frame(input_frame, bg='white')
        self.text_frame.pack(fill='both', expand=True, pady=5)
        
        self.text_area = tk.Text(self.text_frame, height=8, wrap='word', 
                                font=('Arial', 10), bg='#f9f9f9')
        text_scrollbar = tk.Scrollbar(self.text_frame, orient='vertical', command=self.text_area.yview)
        self.text_area.configure(yscrollcommand=text_scrollbar.set)
        
        self.text_area.pack(side='left', fill='both', expand=True)
        text_scrollbar.pack(side='right', fill='y')
        
        # 文件上传框架
        self.file_frame = tk.Frame(input_frame, bg='white')
        
        tk.Button(self.file_frame, text="选择文件", command=self.upload_file, 
                 bg='#4CAF50', fg='white', font=('Arial', 10)).pack(pady=5)
        
        self.file_label = tk.Label(self.file_frame, text="支持 .txt, .docx, .pdf 格式", 
                                  bg='white', fg='#666', font=('Arial', 9))
        self.file_label.pack()
        
        # 参数设置区域
        param_frame = tk.LabelFrame(parent, text="⚙️ 词云参数设置", font=('Arial', 12, 'bold'), 
                                   bg='white', fg='#333', padx=10, pady=10)
        param_frame.pack(fill='x', padx=10, pady=5)
        
        # 尺寸设置
        size_frame = tk.Frame(param_frame, bg='white')
        size_frame.pack(fill='x', pady=5)
        
        tk.Label(size_frame, text="宽度:", bg='white').grid(row=0, column=0, sticky='w', padx=5)
        self.width_var = tk.IntVar(value=800)
        self.width_scale = tk.Scale(size_frame, from_=400, to=1200, orient='horizontal', 
                                   variable=self.width_var, bg='white')
        self.width_scale.grid(row=0, column=1, sticky='ew', padx=5)
        
        tk.Label(size_frame, text="高度:", bg='white').grid(row=1, column=0, sticky='w', padx=5)
        self.height_var = tk.IntVar(value=600)
        self.height_scale = tk.Scale(size_frame, from_=300, to=800, orient='horizontal', 
                                    variable=self.height_var, bg='white')
        self.height_scale.grid(row=1, column=1, sticky='ew', padx=5)
        
        size_frame.columnconfigure(1, weight=1)
        
        # 词汇设置
        word_frame = tk.Frame(param_frame, bg='white')
        word_frame.pack(fill='x', pady=5)
        
        tk.Label(word_frame, text="最大词数:", bg='white').grid(row=0, column=0, sticky='w', padx=5)
        self.max_words_var = tk.IntVar(value=200)
        self.max_words_scale = tk.Scale(word_frame, from_=50, to=500, orient='horizontal', 
                                       variable=self.max_words_var, bg='white')
        self.max_words_scale.grid(row=0, column=1, sticky='ew', padx=5)
        
        tk.Label(word_frame, text="最小字体:", bg='white').grid(row=1, column=0, sticky='w', padx=5)
        self.min_font_var = tk.IntVar(value=10)
        self.min_font_scale = tk.Scale(word_frame, from_=4, to=20, orient='horizontal', 
                                      variable=self.min_font_var, bg='white')
        self.min_font_scale.grid(row=1, column=1, sticky='ew', padx=5)
        
        word_frame.columnconfigure(1, weight=1)
        
        # 颜色设置
        color_frame = tk.Frame(param_frame, bg='white')
        color_frame.pack(fill='x', pady=5)
        
        tk.Label(color_frame, text="颜色主题:", bg='white').pack(anchor='w')
        self.color_theme = tk.StringVar(value="默认")
        theme_combo = ttk.Combobox(color_frame, textvariable=self.color_theme, 
                                  values=["默认", "蓝色系", "红色系", "绿色系", "紫色系", "彩虹色"], 
                                  state='readonly')
        theme_combo.pack(fill='x', pady=2)
        
        # 背景颜色
        bg_color_frame = tk.Frame(color_frame, bg='white')
        bg_color_frame.pack(fill='x', pady=5)
        
        tk.Label(bg_color_frame, text="背景颜色:", bg='white').pack(side='left')
        self.bg_color = "#FFFFFF"
        self.bg_color_btn = tk.Button(bg_color_frame, text="选择颜色", 
                                     command=self.choose_bg_color, bg=self.bg_color, 
                                     width=10)
        self.bg_color_btn.pack(side='right')
        
        # 停用词设置
        stop_frame = tk.LabelFrame(parent, text="🚫 词汇排除设置", font=('Arial', 12, 'bold'), 
                                  bg='white', fg='#333', padx=10, pady=10)
        stop_frame.pack(fill='x', padx=10, pady=5)
        
        tk.Label(stop_frame, text="自定义停用词（用逗号分隔）:", bg='white').pack(anchor='w')
        self.stop_words_var = tk.StringVar(value="的,了,在,是,我,有,和,就,不,人,都,一,一个,上,也,很,到,说,要,去,你,会,着,没有,看,好,自己,这")
        stop_entry = tk.Entry(stop_frame, textvariable=self.stop_words_var, font=('Arial', 9))
        stop_entry.pack(fill='x', pady=5)
        
        # 生成按钮
        generate_btn = tk.Button(parent, text="🎨 生成词云图", command=self.generate_wordcloud, 
                                bg='#2196F3', fg='white', font=('Arial', 14, 'bold'), 
                                height=2)
        generate_btn.pack(fill='x', padx=10, pady=10)
        
    def create_right_panel(self, parent):
        # 显示区域标题
        title_frame = tk.Frame(parent, bg='white')
        title_frame.pack(fill='x', padx=10, pady=5)
        
        tk.Label(title_frame, text="☁️ 词云图显示", font=('Arial', 14, 'bold'), 
                bg='white', fg='#333').pack()
        
        # 词云图显示区域
        self.canvas_frame = tk.Frame(parent, bg='white')
        self.canvas_frame.pack(fill='both', expand=True, padx=10, pady=5)
        
        # 初始提示
        self.info_label = tk.Label(self.canvas_frame, 
                                  text="👈 请在左侧输入文本，然后点击生成按钮\n\n使用说明：\n1. 选择输入方式（直接输入或上传文件）\n2. 输入或上传文本内容\n3. 调整词云参数（可选）\n4. 点击生成词云图按钮\n5. 查看生成的词云图\n6. 可以保存生成的词云图", 
                                  bg='white', fg='#666', font=('Arial', 11), 
                                  justify='left')
        self.info_label.pack(expand=True)
        
        # 按钮区域
        button_frame = tk.Frame(parent, bg='white')
        button_frame.pack(fill='x', padx=10, pady=5)
        
        self.save_btn = tk.Button(button_frame, text="💾 保存词云图", 
                                 command=self.save_wordcloud, bg='#4CAF50', fg='white', 
                                 font=('Arial', 10), state='disabled')
        self.save_btn.pack(side='left', padx=5)
        
        self.stats_btn = tk.Button(button_frame, text="📊 查看词频统计", 
                                  command=self.show_word_stats, bg='#FF9800', fg='white', 
                                  font=('Arial', 10), state='disabled')
        self.stats_btn.pack(side='left', padx=5)
        
    def toggle_input_method(self):
        if self.input_method.get() == "直接输入":
            self.file_frame.pack_forget()
            self.text_frame.pack(fill='both', expand=True, pady=5)
        else:
            self.text_frame.pack_forget()
            self.file_frame.pack(fill='x', pady=5)
            
    def upload_file(self):
        file_path = filedialog.askopenfilename(
            title="选择文本文件",
            filetypes=[("文本文件", "*.txt"), ("Word文档", "*.docx"), ("PDF文件", "*.pdf"), ("所有文件", "*.*")]
        )
        
        if file_path:
            try:
                if file_path.endswith('.txt'):
                    with open(file_path, 'r', encoding='utf-8') as f:
                        self.text_content = f.read()
                elif file_path.endswith('.docx'):
                    try:
                        import docx
                        doc = docx.Document(file_path)
                        self.text_content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                    except ImportError:
                        messagebox.showerror("错误", "请安装 python-docx 库来支持 .docx 文件")
                        return
                elif file_path.endswith('.pdf'):
                    try:
                        import PyPDF2
                        with open(file_path, 'rb') as f:
                            pdf_reader = PyPDF2.PdfReader(f)
                            self.text_content = ""
                            for page in pdf_reader.pages:
                                self.text_content += page.extract_text()
                    except ImportError:
                        messagebox.showerror("错误", "请安装 PyPDF2 库来支持 .pdf 文件")
                        return
                else:
                    # 尝试以文本方式读取
                    with open(file_path, 'r', encoding='utf-8') as f:
                        self.text_content = f.read()
                        
                self.file_label.config(text=f"已选择: {os.path.basename(file_path)}")
                messagebox.showinfo("成功", "文件读取成功！")
                
            except Exception as e:
                messagebox.showerror("错误", f"文件读取错误: {str(e)}")
                
    def choose_bg_color(self):
        color = colorchooser.askcolor(title="选择背景颜色", initialcolor=self.bg_color)
        if color[1]:  # 如果用户选择了颜色
            self.bg_color = color[1]
            self.bg_color_btn.config(bg=self.bg_color)
            
    def generate_wordcloud(self):
        # 获取文本内容
        if self.input_method.get() == "直接输入":
            self.text_content = self.text_area.get("1.0", tk.END).strip()
        
        if not self.text_content.strip():
            messagebox.showwarning("警告", "请先输入文本内容")
            return
            
        try:
            # 清除之前的显示
            for widget in self.canvas_frame.winfo_children():
                widget.destroy()
                
            # 显示进度提示
            progress_label = tk.Label(self.canvas_frame, text="正在生成词云图，请稍候...", 
                                    bg='white', fg='#666', font=('Arial', 12))
            progress_label.pack(expand=True)
            self.root.update()
            
            # 文本预处理
            text_clean = re.sub(r'[^\u4e00-\u9fa5a-zA-Z\s]', '', self.text_content)
            
            # 中文分词
            words = jieba.lcut(text_clean)
            
            # 处理停用词
            user_stop_words = set([word.strip() for word in self.stop_words_var.get().split(',') if word.strip()])
            
            # 过滤停用词和短词
            words = [word for word in words if len(word) > 1 and word not in user_stop_words]
            
            if not words:
                messagebox.showerror("错误", "没有找到有效的词语，请检查输入文本")
                return
                
            # 词频统计
            self.word_freq = Counter(words)
            
            # 设置颜色映射
            color_maps = {
                "默认": None,
                "蓝色系": "Blues",
                "红色系": "Reds",
                "绿色系": "Greens",
                "紫色系": "Purples",
                "彩虹色": "rainbow"
            }
            
            # 生成词云
            self.current_wordcloud = WordCloud(
                font_path='simhei.ttf',  # 中文字体
                width=self.width_var.get(),
                height=self.height_var.get(),
                background_color=self.bg_color,
                max_words=self.max_words_var.get(),
                min_font_size=self.min_font_var.get(),
                colormap=color_maps[self.color_theme.get()],
                relative_scaling=0.5,
                random_state=42
            ).generate_from_frequencies(self.word_freq)
            
            # 清除进度提示
            progress_label.destroy()
            
            # 显示词云图
            self.current_figure, ax = plt.subplots(figsize=(10, 6))
            ax.imshow(self.current_wordcloud, interpolation='bilinear')
            ax.axis('off')
            
            # 在Tkinter中显示matplotlib图形
            canvas = FigureCanvasTkAgg(self.current_figure, self.canvas_frame)
            canvas.draw()
            canvas.get_tk_widget().pack(fill='both', expand=True)
            
            # 启用按钮
            self.save_btn.config(state='normal')
            self.stats_btn.config(state='normal')
            
            messagebox.showinfo("成功", "词云图生成完成！")
            
        except Exception as e:
            messagebox.showerror("错误", f"生成词云图时出错: {str(e)}\n\n提示：如果遇到字体问题，请确保系统中有中文字体文件")
            
    def save_wordcloud(self):
        if self.current_figure is None:
            messagebox.showwarning("警告", "请先生成词云图")
            return
            
        file_path = filedialog.asksaveasfilename(
            title="保存词云图",
            defaultextension=".png",
            filetypes=[("PNG图片", "*.png"), ("JPEG图片", "*.jpg"), ("所有文件", "*.*")]
        )
        
        if file_path:
            try:
                self.current_figure.savefig(file_path, bbox_inches='tight', dpi=300)
                messagebox.showinfo("成功", f"词云图已保存到: {file_path}")
            except Exception as e:
                messagebox.showerror("错误", f"保存失败: {str(e)}")
                
    def show_word_stats(self):
        if not hasattr(self, 'word_freq') or not self.word_freq:
            messagebox.showwarning("警告", "请先生成词云图")
            return
            
        # 创建词频统计窗口
        stats_window = tk.Toplevel(self.root)
        stats_window.title("词频统计")
        stats_window.geometry("400x600")
        stats_window.configure(bg='white')
        
        # 标题
        tk.Label(stats_window, text="📊 词频统计 (前30个)", 
                font=('Arial', 14, 'bold'), bg='white').pack(pady=10)
        
        # 创建表格
        frame = tk.Frame(stats_window, bg='white')
        frame.pack(fill='both', expand=True, padx=10, pady=5)
        
        # 表格标题
        header_frame = tk.Frame(frame, bg='#f0f0f0')
        header_frame.pack(fill='x')
        
        tk.Label(header_frame, text="排名", width=8, bg='#f0f0f0', 
                font=('Arial', 10, 'bold')).pack(side='left')
        tk.Label(header_frame, text="词语", width=15, bg='#f0f0f0', 
                font=('Arial', 10, 'bold')).pack(side='left')
        tk.Label(header_frame, text="频次", width=8, bg='#f0f0f0', 
                font=('Arial', 10, 'bold')).pack(side='left')
        
        # 滚动框架
        canvas = tk.Canvas(frame, bg='white')
        scrollbar = tk.Scrollbar(frame, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas, bg='white')
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        # 添加词频数据
        top_words = self.word_freq.most_common(30)
        for i, (word, freq) in enumerate(top_words, 1):
            row_frame = tk.Frame(scrollable_frame, bg='white')
            row_frame.pack(fill='x', pady=1)
            
            tk.Label(row_frame, text=str(i), width=8, bg='white').pack(side='left')
            tk.Label(row_frame, text=word, width=15, bg='white').pack(side='left')
            tk.Label(row_frame, text=str(freq), width=8, bg='white').pack(side='left')
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # 关闭按钮
        tk.Button(stats_window, text="关闭", command=stats_window.destroy, 
                 bg='#f44336', fg='white', font=('Arial', 10)).pack(pady=10)

def main():
    root = tk.Tk()
    app = WordCloudGenerator(root)
    root.mainloop()

if __name__ == "__main__":
    main()
